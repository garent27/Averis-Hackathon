#!/usr/bin/env python3
"""
extract_compare.py — extract the 7 shipment fields from SI/BL attachments
(.txt/.pdf/.docx/.xlsx) and compare them.

Hybrid extraction, cheapest/most-reliable tier first:
  1. Deterministic label-matching (regex/synonym dictionary, copied from
     sdoc-docker/data_v2/pools.py LABELS) -- free, exact, and very reliable
     on this dataset's structured layouts. Handles the "same information
     looks different" synonym challenge from the spec (e.g. "Port of
     Loading" vs "Load Port").
  2. If step 1 didn't find all 7 fields, send the attachment's plain text to
     the LLM extractor (llm.extract_fields_raw) to fill in ONLY the missing
     fields -- deterministic values are trusted over LLM ones wherever both
     exist.
  3. If there's no usable text at all (image-only/scanned PDF -- the
     "Scanned documents" advanced-stage case), render the first page to an
     image and use the vision LLM extractor (llm.extract_fields_vision).

Layout per format, confirmed by inspecting real generated samples
(sdoc-docker/data_v2/render.py + attachments/email_055_BL.docx,
email_055_SI.xlsx, email_059_BL.pdf):

  .txt  - one "Label: value" line per field (data_v2/render.py si_txt/bl_txt).
  .docx - a 2-column table; col0 = "<label> (<chinese annotation>)",
          col1 = "<value>\\n<address line 1>\\n..." for name fields.
  .xlsx - 2-column rows; col0 = exact label text, col1 = "<name> | <addr1>;
          <addr2>" for name fields, a plain number for gross_weight_kg.
  .pdf  - plain extracted text. Name/port fields render as "<label> <value>"
          on one line with NO colon (block() helper draws label/value at
          different x positions, pdfplumber's extract_text() just joins
          them with a space); container_count/gross_weight_kg DO use
          "<label>: <value>" (optionally "TOTAL " prefixed for weight).

Then step 3's values are normalized (case, whitespace, punctuation, company
suffixes, weight units) into a canonical form before comparison.
"""
import io
import re

COMPARE_FIELDS = [
    "shipper", "consignee", "notify_party",
    "port_of_loading", "port_of_discharge",
    "container_count", "gross_weight_kg",
]
NAME_FIELDS = {"shipper", "consignee", "notify_party"}

# Trailing company-suffix tokens to ignore when comparing names -- "Vital
# Solutions Pte Ltd" and "Vital Solutions Pte Ltd." should compare equal even
# if an LLM extraction is slightly inconsistent about trailing punctuation.
_COMPANY_SUFFIXES = {"LTD", "LIMITED", "LLC", "INC", "CORP", "CO", "COMPANY", "GMBH", "FZE"}

# Canonical field -> label variants actually used by data_v2/render.py.
FIELD_LABELS = {
    "shipper": ["Shipper", "Shipper/Exporter", "Shipper (Principal or Seller)", "SHIPPER"],
    "consignee": ["Consignee", "Consignee (Non-Negotiable)", "CONSIGNEE", "To the Order of"],
    "notify_party": ["Notify Party", "Notify", "Notify Party/Intermediate Consignee", "NOTIFY PARTY"],
    "port_of_loading": ["Port of Loading", "Port of Loading (POL)", "Load Port", "POL", "PORT OF LOADING"],
    "port_of_discharge": ["Port of Discharge", "Port of Discharge (POD)", "Discharge Port", "POD", "PORT OF DISCHARGE"],
    "container_count": ["No. of Containers", "Total Containers", "No. of Containers or Packages", "Container Count"],
    "gross_weight_kg": ["Gross Weight (KG)", "Gross Wt (kgs)", "Gross Weight毛重(KGS)", "GROSS WEIGHT"],
}

# label text (uppercased) -> canonical field, for exact-match lookup.
_LABEL_TO_FIELD = {
    variant.upper(): field
    for field, variants in FIELD_LABELS.items()
    for variant in variants
}

# (variant_upper, field), longest variant first so e.g. "Port of Loading
# (POL)" is tried before the shorter "POL" when prefix-matching PDF lines.
_VARIANTS_BY_LENGTH = sorted(
    ((variant.upper(), field) for field, variants in FIELD_LABELS.items() for variant in variants),
    key=lambda pair: len(pair[0]),
    reverse=True,
)


# ---------------------------------------------------------------------------
# shared value normalisation
# ---------------------------------------------------------------------------
def _normalize_port(value: str) -> str:
    # Some renderings append " (CODE)" to the port name; strip it so
    # "Singapore, Singapore (SGSIN)" and "Singapore, Singapore" compare equal.
    value = re.sub(r"\s*\([A-Z0-9]+\)\s*$", "", value.strip())
    return re.sub(r"\s+", " ", value).strip().upper()


def _normalize_container_count(value):
    # "{count} x {size}" in every format, e.g. "3 x 40'HC" -- also tolerates
    # a bare number.
    m = re.match(r"\s*(\d+)", str(value))
    return int(m.group(1)) if m else None


def _normalize_gross_weight(value):
    """Parse a weight value + unit into an integer number of kilograms.
    Handles kg (default), MT/metric tons, and lbs/pounds."""
    if isinstance(value, (int, float)):
        return int(value)
    s = str(value).upper()
    digits = re.sub(r"[^\d.]", "", s)
    if not digits:
        return None
    amount = float(digits)
    if "MT" in s or "TONNE" in s or "METRIC TON" in s:
        amount *= 1000
    elif "LB" in s or "POUND" in s:
        amount *= 0.45359237
    # else: assume kilograms (KG/KGS/no unit)
    return int(round(amount))


def _normalize_name(value: str) -> str:
    v = re.sub(r"[.,]", "", value.upper())
    v = re.sub(r"\s+", " ", v).strip()
    words = v.split()
    while words and words[-1] in _COMPANY_SUFFIXES:
        words.pop()
    return " ".join(words)


def _normalize_value(field, value):
    if value is None:
        return None
    if field == "port_of_loading" or field == "port_of_discharge":
        return _normalize_port(str(value))
    if field == "container_count":
        return _normalize_container_count(value)
    if field == "gross_weight_kg":
        return _normalize_gross_weight(value)
    if field in NAME_FIELDS:
        return _normalize_name(str(value))
    return str(value).strip()


def _finalize(raw: dict) -> dict:
    return {f: v for f in COMPARE_FIELDS if (v := _normalize_value(f, raw.get(f))) is not None}


# ---------------------------------------------------------------------------
# tier 1: deterministic label-matching, per format
# ---------------------------------------------------------------------------
def _match_colon_label(label_text: str):
    """Match a 'Label:' fragment to a canonical field, tolerating a leading
    'TOTAL ' (the PDF's weight line is 'TOTAL Gross Weight (KG): ...')."""
    label = label_text.strip().upper()
    if label.startswith("TOTAL "):
        label = label[len("TOTAL "):].strip()
    return _LABEL_TO_FIELD.get(label)


def _match_prefix_label(line_upper: str):
    for variant_upper, field in _VARIANTS_BY_LENGTH:
        if line_upper.startswith(variant_upper + " "):
            return field, len(variant_upper)
    return None, None


def _extract_raw_txt(text: str) -> dict:
    raw = {}
    for line in text.splitlines():
        if ":" not in line:
            continue
        label, _, value = line.partition(":")
        field = _match_colon_label(label)
        if field is None:
            continue
        value = value.strip()
        raw[field] = value.split("\n")[0].strip() if field in NAME_FIELDS else value
    return raw


def _extract_raw_pdf_text(text: str) -> dict:
    """Label-match a PDF's already-extracted text (colon lines for
    count/weight, label-prefixed lines with no colon for names/ports)."""
    raw = {}
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        if ":" in line:
            label, _, value = line.partition(":")
            field = _match_colon_label(label)
            if field:
                raw[field] = value.strip()
                continue
        field, prefix_len = _match_prefix_label(line.upper())
        if field:
            raw[field] = line[prefix_len:].strip()
    return raw


def _extract_raw_docx(raw_bytes: bytes) -> dict:
    import docx
    d = docx.Document(io.BytesIO(raw_bytes))
    raw = {}
    for table in d.tables:
        for row in table.rows:
            if len(row.cells) < 2:
                continue
            label_upper = row.cells[0].text.strip().upper()
            field, _ = _match_prefix_label(label_upper + " ")  # reuse prefix matcher
            if field is None:
                field = _LABEL_TO_FIELD.get(label_upper)  # exact match fallback
            if field is None:
                continue
            value = row.cells[1].text.strip()
            raw[field] = value.split("\n")[0].strip() if field in NAME_FIELDS else value
    return raw


def _extract_raw_xlsx(raw_bytes: bytes) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
    ws = wb.active
    raw = {}
    for row in ws.iter_rows():
        if len(row) < 2 or row[0].value is None or row[1].value is None:
            continue
        field = _LABEL_TO_FIELD.get(str(row[0].value).strip().upper())
        if field is None:
            continue
        value = row[1].value
        if field in NAME_FIELDS:
            value = str(value).split(" | ")[0].strip()
        raw[field] = value
    return raw


# ---------------------------------------------------------------------------
# plain-text access, shared by tier 1 (pdf) and tier 2 (LLM text fallback)
# ---------------------------------------------------------------------------
def _pdf_text(raw_bytes: bytes) -> str:
    import pdfplumber
    with pdfplumber.open(io.BytesIO(raw_bytes)) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def _docx_text(raw_bytes: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(raw_bytes))
    lines = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def _xlsx_text(raw_bytes: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
    ws = wb.active
    lines = []
    for row in ws.iter_rows():
        values = [str(cell.value) for cell in row if cell.value is not None]
        if values:
            lines.append(" | ".join(values))
    return "\n".join(lines)


def _raw_document_text(raw_bytes: bytes, filename: str) -> str:
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext == "txt":
        return raw_bytes.decode("utf-8", errors="replace")
    if ext == "pdf":
        return _pdf_text(raw_bytes)
    if ext == "docx":
        return _docx_text(raw_bytes)
    if ext == "xlsx":
        return _xlsx_text(raw_bytes)
    return ""


def _render_pdf_first_page_png(raw_bytes: bytes) -> bytes:
    """Render page 1 of a (likely image-only/scanned) PDF to PNG bytes, for
    the vision-LLM fallback tier."""
    import pdfplumber
    with pdfplumber.open(io.BytesIO(raw_bytes)) as pdf:
        image = pdf.pages[0].to_image(resolution=150)
        buf = io.BytesIO()
        image.original.save(buf, format="PNG")
        return buf.getvalue()


# ---------------------------------------------------------------------------
# public API: hybrid extract (tier 1 rules -> tier 2 LLM text -> tier 3 vision)
# ---------------------------------------------------------------------------
def extract_fields(raw_bytes: bytes, filename: str) -> dict:
    """Extract+normalise the 7 comparison fields from an SI/BL attachment,
    hybrid style: deterministic label-matching first (free, exact), LLM text
    fallback only for fields it couldn't find, vision LLM fallback only when
    there's no text at all (scanned/image-only PDF). Missing fields are
    simply absent from the returned dict -- the caller decides what to do
    about an incomplete extraction (see HACKATHON_CONTEXT.md: that's a
    NEEDS_REVIEW case, not a guessed MISMATCH)."""
    ext = filename.rsplit(".", 1)[-1].lower()

    # tier 1: deterministic label-matching
    text = ""
    if ext == "txt":
        text = raw_bytes.decode("utf-8", errors="replace")
        raw = _extract_raw_txt(text)
    elif ext == "pdf":
        text = _pdf_text(raw_bytes)
        raw = _extract_raw_pdf_text(text)
    elif ext == "docx":
        raw = _extract_raw_docx(raw_bytes)
    elif ext == "xlsx":
        raw = _extract_raw_xlsx(raw_bytes)
    else:
        raw = {}

    missing = [f for f in COMPARE_FIELDS if f not in raw]
    if not missing:
        return _finalize(raw)

    # tier 2: LLM text fallback, only to fill what tier 1 missed. A fallback
    # failure (no API key, network error, etc.) must not lose tier 1's
    # results -- fall through with whatever tier 1 already found rather than
    # raising, so one missing field doesn't wipe the other six.
    import llm  # imported lazily so this module works without an API key when tier 1 suffices

    if ext != "pdf":
        text = _raw_document_text(raw_bytes, filename)
    if text.strip():
        try:
            llm_fields = llm.extract_fields_raw(text)
        except Exception as e:
            print(f"  ! LLM text fallback failed ({filename}): {e}")
            return _finalize(raw)
        for field in missing:
            if field in llm_fields:
                raw[field] = llm_fields[field]
        return _finalize(raw)

    # tier 3: no text layer at all (image-only/scanned PDF) -- vision fallback
    if ext == "pdf":
        try:
            image_bytes = _render_pdf_first_page_png(raw_bytes)
            vision_fields = llm.extract_fields_vision(image_bytes)
        except Exception as e:
            print(f"  ! vision fallback failed ({filename}): {e}")
            return _finalize(raw)
        for field in missing:
            if field in vision_fields:
                raw[field] = vision_fields[field]

    return _finalize(raw)


def compare_fields(si_fields: dict, bl_fields: dict) -> list:
    """Return the sorted list of the 7 fields whose SI and BL values differ.
    A field missing from either side is NOT reported as a mismatch -- a
    blank/unreadable value is uncertainty, not a discrepancy (see
    HACKATHON_CONTEXT.md: NEEDS_REVIEW is a separate axis from MISMATCH)."""
    mismatches = []
    for field in COMPARE_FIELDS:
        si_val = si_fields.get(field)
        bl_val = bl_fields.get(field)
        if si_val is None or bl_val is None:
            continue
        if si_val != bl_val:
            mismatches.append(field)
    return sorted(mismatches)
